from docx import Document
from PIL import Image
import os
import chardet


def extract_images_and_create_new_docx(source_file_path, output_folder):
    # Step 1: Read the source DOCX file
    doc = Document(source_file_path)

    # Step 2: Extract images from the source document
    image_folder = os.path.join(output_folder, 'images')
    os.makedirs(image_folder, exist_ok=True)

    images = {}  # A dictionary to store image placeholders and their actual paths

    for rel in doc.part.rels:
        if 'image' in doc.part.rels[rel].target_ref:
            image_part = doc.part.rels[rel].target_part
            image_ext = image_part.content_type.split('/')[-1]
            image_name = f"image_{rel}.{image_ext}"
            image_path = os.path.join(image_folder, image_name)

            with open(image_path, 'wb') as image_file:
                image_file.write(image_part.blob)

            images[rel] = image_path

    # Step 3: Replace image placeholders in the document content
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for rel, image_path in images.items():
                placeholder = f"IMAGE_SOURCEWILLBEHERE:{rel}"
                run.text = run.text.replace(placeholder, image_path)

    # Step 4: Save the new DOCX document with replaced image placeholders
    new_docx_path = os.path.join(output_folder, 'new_doc.docx')
    doc.save(new_docx_path)
    return new_docx_path

# Example usage
source_file_path = "E:\Project\Job\GQ\Python\Baidu_downlaod_upload\Baidu_Pyqt5_Threading\document_merger\docx_formatinng\ok.docx"
output_folder = "E:\Project\Job\GQ\Python\Baidu_downlaod_upload\Baidu_Pyqt5_Threading\document_merger\docx_formatinng\\test"
new_docx_path = extract_images_and_create_new_docx(source_file_path, output_folder)
if new_docx_path:
    print(f"Formatting completed. New DOCX saved at: {new_docx_path}")
else:
    print("Formatting failed.")