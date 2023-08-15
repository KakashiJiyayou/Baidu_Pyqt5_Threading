from docx import Document
from docx.oxml import OxmlElement
import os

def extract_images_and_create_new_docx(source_file_path, output_folder, index):
    # Step 1: Read the source DOCX file
    doc = Document(source_file_path)

    # Step 2: Extract images from the source document
    image_folder = os.path.join(output_folder, 'images')
    os.makedirs(image_folder, exist_ok=True)

    images = {}  # A dictionary to store image placeholders and their actual paths

    for rel in doc.part.rels:
        if 'image' in doc.part.rels[rel].target_ref:
            image_part = doc.part.rels[rel].target_part
            image_ext = image_part.content_type.split('/')[-1]
            image_name = f"image_{rel}.{image_ext}"
            image_path = os.path.join(image_folder, image_name)

            with open(image_path, 'wb') as image_file:
                image_file.write(image_part.blob)

            images[rel] = image_path

    # Step 3: Replace image placeholders in the document content
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for rel, image_path in images.items():
                placeholder = f"IMAGE_SOURCEWILLBEHERE:{rel}"

                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, "")
                    run_element = run._r
                    drawing = OxmlElement('w:drawing')
                    inline = OxmlElement('wp:inline')
                    extent = OxmlElement('wp:extent')
                    extent.set('cx', '4000000')
                    extent.set('cy', '3000000')
                    effect_extent = OxmlElement('wp:effectExtent')
                    effect_extent.set('l', '0')
                    effect_extent.set('t', '0')
                    effect_extent.set('r', '0')
                    effect_extent.set('b', '0')
                    doc_pr = OxmlElement('wp:docPr')
                    doc_pr.set('id', '1')
                    doc_pr.set('name', 'Image 1')
                    c_nv_pr = OxmlElement('wp:cNvPr')
                    c_nv_pr.set('id', '2')
                    c_nv_pr.set('name', 'Picture 1')
                    c_nv_pr.set('descr', image_path)
                    blip_fill = OxmlElement('a:blip')
                    blip_fill.set('r:embed', rel)
                    stretch = OxmlElement('a:stretch')
                    fill_rect = OxmlElement('a:fillRect')
                    stretch.append(fill_rect)
                    blip_fill.append(stretch)
                    blip = OxmlElement('a:blipFill')
                    blip.append(blip_fill)
                    src_rect = OxmlElement('a:srcRect')
                    src_rect.set('t', '0')
                    src_rect.set('r', '0')
                    src_rect.set('b', '0')
                    src_rect.set('l', '0')
                    stretch_info = OxmlElement('a:stretch')
                    stretch_info.append(src_rect)
                    blip.append(stretch_info)
                    inline.append(extent)
                    inline.append(effect_extent)
                    inline.append(doc_pr)
                    inline.append(c_nv_pr)
                    inline.append(blip)
                    drawing.append(inline)
                    run_element.append(drawing)

    # Step 4: Save the new DOCX document with replaced image placeholders
    new_docx_name = f"tempx{index}.docx"  # Change the naming convention as needed
    new_docx_path = os.path.join(output_folder, new_docx_name)
    doc.save(new_docx_path)

    # Step 5: Clean up image folder after creating new docx
    for image_path in images.values():
        os.remove(image_path)  # Delete individual images

    os.rmdir(image_folder)  # Delete the image folder

    return new_docx_path

# Example usage
source_file_path = "E:\Project\Job\GQ\Python\Baidu_downlaod_upload\Baidu_Pyqt5_Threading\document_merger\docx_formatinng\doc3.docx"
output_folder = "E:\Project\Job\GQ\Python\Baidu_downlaod_upload\Baidu_Pyqt5_Threading\document_merger\docx_formatinng\\test"
index = 4  # Adjust the index as needed
new_docx_path = extract_images_and_create_new_docx(source_file_path, output_folder, index)
if new_docx_path:
    print(f"Formatting completed. New DOCX saved at: {new_docx_path}")
else:
    print("Formatting failed.")
